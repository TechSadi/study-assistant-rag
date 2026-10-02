"""Read an uploaded document into pages of text for rag.split_text.

Every format becomes the same records, one per page with text:
{text, source, page (1-based)}, where text is paragraphs separated by blank lines.
"Page" means what a student would cite:

* PDF, EPUB   the document's pages (PyMuPDF)
* PowerPoint  one page per slide, with the speaker notes
* Word        Word's own page breaks as last saved; if the file has none
              (e.g. made by Google Docs), sections of about TEXT_PAGE_CHARS
* Text, MD    sections of about TEXT_PAGE_CHARS, at paragraph boundaries
"""

import re
from pathlib import Path

import pymupdf

from . import config

# Extension -> the name shown in messages. The frontend keeps the same list (App.jsx).
FILE_TYPES = {
    ".pdf": "PDF",
    ".docx": "Word",
    ".pptx": "PowerPoint",
    ".epub": "EPUB",
    ".txt": "text",
    ".md": "Markdown",
    ".markdown": "Markdown",
}
# Older binary Office formats that can't be read: suggest converting them.
OLD_OFFICE_TYPES = {".doc": ".docx", ".ppt": ".pptx"}

_ZIP_TYPES = {".docx", ".pptx", ".epub"}
_TEXT_TYPES = {".txt", ".md", ".markdown"}


def file_kind(filename):
    """The FILE_TYPES name for a filename, or None if it can't be uploaded."""
    return FILE_TYPES.get(Path(filename).suffix.lower())


def looks_valid(filename, data):
    """A quick check that the bytes match the extension, before anything is saved."""
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        return data.startswith(b"%PDF-")
    if suffix in _ZIP_TYPES:
        return data.startswith(b"PK\x03\x04")
    if suffix in _TEXT_TYPES:
        # UTF-16 text has NUL bytes, anything else with them is binary.
        return data.startswith((b"\xff\xfe", b"\xfe\xff")) or b"\x00" not in data[:8192]
    return False


def extract_pages(path, source_name):
    """Return (page_count, pages), one record per page with text: {text, source, page}.

    Raises ValueError with a readable reason when the file can't be used.
    Scanned or image-only PDFs have no pages with text because there is no OCR.
    """
    suffix = Path(path).suffix.lower()
    if suffix in (".pdf", ".epub"):
        return _pymupdf_pages(path, source_name)
    if suffix == ".docx":
        texts = _docx_pages(path)
    elif suffix == ".pptx":
        texts = _pptx_pages(path)
    elif suffix in _TEXT_TYPES:
        text = _read_text(path)
        texts = _paginate(_markdown_text(text) if suffix != ".txt" else _plain_text(text))
    else:
        raise ValueError(f"{suffix or 'This'} files aren't supported.")
    pages = [
        {"text": text, "source": source_name, "page": number}
        for number, text in enumerate(texts, start=1)
        if text
    ]
    return len(texts), pages


# --- PDF and EPUB ------------------------------------------------------------

def page_text(page):
    """A page's text as paragraphs separated by blank lines.

    PyMuPDF's text blocks are roughly paragraphs (a heading, a paragraph, a table
    cell). The line breaks inside a block are only where the line wrapped, so they
    are joined with spaces, and words hyphenated across lines are rejoined.
    """
    blocks = page.get_text("blocks", flags=pymupdf.TEXTFLAGS_BLOCKS | pymupdf.TEXT_DEHYPHENATE)
    paragraphs = (" ".join(b[4].split()) for b in blocks if b[6] == 0)  # type 0 = text
    return "\n\n".join(p for p in paragraphs if p)


def _pymupdf_pages(path, source_name):
    pages = []
    with pymupdf.open(path) as doc:
        if doc.needs_pass:
            raise ValueError("This file is password-protected.")
        for page_num, page in enumerate(doc):
            text = page_text(page)
            if text:
                pages.append({"text": text, "source": source_name, "page": page_num + 1})
        return doc.page_count, pages


# --- Word ----------------------------------------------------------------------

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _docx_pages(path):
    import docx
    from docx.table import Table

    document = docx.Document(str(path))
    body = document.element.body
    # Word records where it broke pages when it last saved the file. Without those,
    # the document is paginated by length, still honouring manual page breaks.
    rendered = next(body.iter(f"{_W}lastRenderedPageBreak"), None) is not None

    pages, current = [], []
    for block in document.iter_inner_content():
        if isinstance(block, Table):
            text = _docx_table(block)
            breaks = 0
        else:
            text = _docx_paragraph(block)
            breaks = _docx_breaks(block._p, rendered)
        # A page that starts inside this paragraph gets the paragraph.
        for _ in range(breaks):
            pages.append(current)
            current = []
        if text:
            current.append(text)
    pages.append(current)

    if rendered:
        return ["\n\n".join(p) for p in pages]
    return [page for paragraphs in pages for page in _paginate("\n\n".join(paragraphs)) if page]


def _docx_breaks(p, rendered):
    """How many page breaks are in a paragraph (or before it)."""
    count = sum(1 for br in p.iter(f"{_W}br") if br.get(f"{_W}type") == "page")
    if p.find(f"{_W}pPr/{_W}pageBreakBefore") is not None:
        count += 1
    if rendered:
        count += sum(1 for _ in p.iter(f"{_W}lastRenderedPageBreak"))
    return count


def _docx_paragraph(paragraph):
    text = " ".join(paragraph.text.split())
    style = paragraph.style.name if paragraph.style is not None else ""
    if text and style.startswith("List"):
        return "• " + text
    return text


def _docx_table(table):
    """One paragraph per row, cells separated by " | " (merged cells only once)."""
    rows = []
    for row in table.rows:
        cells = []
        for cell in row.cells:
            text = " ".join(cell.text.split())
            if text and (not cells or cells[-1] != text):
                cells.append(text)
        if cells:
            rows.append(" | ".join(cells))
    return "\n\n".join(rows)


# --- PowerPoint ------------------------------------------------------------------

def _pptx_pages(path):
    from pptx import Presentation

    pages = []
    for slide in Presentation(str(path)).slides:
        paragraphs = []
        for shape in slide.shapes:
            paragraphs.extend(_pptx_shape(shape))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            if notes is not None and notes.text.strip():
                paragraphs.append("Speaker notes: " + " ".join(notes.text.split()))
        pages.append("\n\n".join(paragraphs))
    return pages


def _pptx_shape(shape):
    if shape.shape_type == 6:  # a group: its shapes in order
        return [p for child in shape.shapes for p in _pptx_shape(child)]
    if getattr(shape, "has_table", False) and shape.has_table:
        rows = (" | ".join(" ".join(c.text.split()) for c in row.cells if c.text.strip()) for row in shape.table.rows)
        return [r for r in rows if r]
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        lines = (" ".join(p.text.split()) for p in shape.text_frame.paragraphs)
        return [line for line in lines if line]
    return []


# --- Text and Markdown -----------------------------------------------------------

def _read_text(path):
    data = Path(path).read_bytes()
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


def _plain_text(text):
    return text.replace("\r\n", "\n").replace("\r", "\n")


_MD_HEADING = re.compile(r"^\s{0,3}#{1,6}\s+(.*?)\s*#*\s*$")
_MD_LIST = re.compile(r"^\s*(?:[-*+]|\d+[.)])\s+")
_MD_FENCE = re.compile(r"^\s*(```|~~~)")
_MD_RULE = re.compile(r"^\s{0,3}([-*_])(\s*\1){2,}\s*$")
_MD_TABLE_DIVIDER = re.compile(r"^\s*\|?\s*:?-+:?\s*(\|\s*:?-+:?\s*)*\|?\s*$")


def _markdown_inline(line):
    line = re.sub(r"!\[([^\]]*)\]\([^)]*\)", r"\1", line)  # images -> alt text
    line = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", line)   # links -> link text
    line = re.sub(r"(\*\*|__)(.+?)\1", r"\2", line)          # bold
    line = re.sub(r"(?<![\w*])\*(?!\s)(.+?)(?<!\s)\*(?![\w*])", r"\1", line)  # italics
    return line


def _markdown_text(text):
    """Markdown as plain paragraphs: headings, list items and table rows each become
    a paragraph of their own, and formatting marks are removed. Code is kept."""
    text = _plain_text(text)
    text = re.sub(r"\A---\n.*?\n(---|\.\.\.)\n", "", text, flags=re.S)  # front matter
    text = re.sub(r"<!--.*?-->", "", text, flags=re.S)

    paragraphs, current, in_code = [], [], False

    def flush():
        if current:
            paragraphs.append(" ".join(current))
            current.clear()

    for line in text.split("\n"):
        if _MD_FENCE.match(line):
            flush()
            in_code = not in_code
            continue
        if in_code:
            if line.strip():
                paragraphs.append(line.strip())  # a line of code
            continue
        stripped = line.strip()
        if not stripped or _MD_RULE.match(line) or _MD_TABLE_DIVIDER.match(line):
            flush()
            continue
        heading = _MD_HEADING.match(line)
        if heading:
            flush()
            paragraphs.append(_markdown_inline(heading.group(1)))
        elif _MD_LIST.match(line) or stripped.startswith("|"):
            flush()
            current.append(_markdown_inline(_MD_LIST.sub("• ", line, count=1).strip().strip("|").strip()))
        else:
            current.append(_markdown_inline(stripped.lstrip("> ").strip()))
    flush()
    return "\n\n".join(p for p in paragraphs if p.strip())


def _paginate(text, size=None):
    """Group paragraphs into pages of about `size` characters (a printed page).
    A paragraph is never split across pages."""
    size = size or config.TEXT_PAGE_CHARS
    pages, current, length = [], [], 0
    for paragraph in re.split(r"\n\s*\n", text):
        paragraph = " ".join(paragraph.split())
        if not paragraph:
            continue
        if current and length + len(paragraph) > size:
            pages.append("\n\n".join(current))
            current, length = [], 0
        current.append(paragraph)
        length += len(paragraph) + 2
    if current:
        pages.append("\n\n".join(current))
    return pages
