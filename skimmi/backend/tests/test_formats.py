"""Uploading documents other than PDFs: each becomes pages of text with citations."""

import io

from app import config, extract, rag

from .conftest import USER_A, upload, stored_path


def make_docx(paragraphs, page_break_after=None):
    import docx

    document = docx.Document()
    document.add_heading("Cell Biology", level=1)
    for i, text in enumerate(paragraphs):
        document.add_paragraph(text)
        if i == page_break_after:
            document.add_page_break()
    table = document.add_table(rows=2, cols=2)
    for r, row in enumerate([("Organelle", "Role"), ("Ribosome", "Makes proteins")]):
        for c, text in enumerate(row):
            table.cell(r, c).text = text
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def make_pptx(slides):
    from pptx import Presentation

    deck = Presentation()
    for title, body, notes in slides:
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
    out = io.BytesIO()
    deck.save(out)
    return out.getvalue()


def test_docx_pages_follow_manual_page_breaks(client):
    data = make_docx(["Cells are the unit of life.", "Mitochondria make ATP."], page_break_after=0)
    doc = upload(client, "biology.docx", data).json()
    assert doc["filename"] == "biology.docx" and doc["pages"] == 2
    assert (stored_path(USER_A, doc["id"], ".docx")).exists()

    _, pages = extract.extract_pages(stored_path(USER_A, doc["id"], ".docx"), "biology.docx")
    assert pages[0]["text"] == "Cell Biology\n\nCells are the unit of life."
    assert pages[1]["text"] == "Mitochondria make ATP.\n\nOrganelle | Role\n\nRibosome | Makes proteins"

    sources = client.post("/api/chat", json={"question": "What makes ATP?"}).json()["message"]["sources"]
    assert {(s["source"], s["page"]) for s in sources} == {("biology.docx", 1), ("biology.docx", 2)}

    assert client.delete(f"/api/documents/{doc['id']}").status_code == 204
    assert not config.UPLOAD_DIR.exists() or not any(p.is_file() for p in config.UPLOAD_DIR.rglob('*'))


def test_pptx_has_one_page_per_slide_with_notes(client):
    data = make_pptx([("Photosynthesis", "Light becomes chemical energy", "Mention chlorophyll."), ("Respiration", "Glucose is broken down", "")])
    doc = upload(client, "lecture.pptx", data).json()
    assert doc["pages"] == 2 and doc["chunks"] == 2

    _, pages = extract.extract_pages(stored_path(USER_A, doc["id"], ".pptx"), "lecture.pptx")
    assert pages[0]["text"] == "Photosynthesis\n\nLight becomes chemical energy\n\nSpeaker notes: Mention chlorophyll."
    assert pages[1] == {"text": "Respiration\n\nGlucose is broken down", "source": "lecture.pptx", "page": 2}


def test_text_files_are_cut_into_pages(client, monkeypatch):
    monkeypatch.setattr(config, "TEXT_PAGE_CHARS", 60)
    text = "First paragraph about alpha.\r\n\r\nSecond paragraph about beta.\n\nThird paragraph about gamma."
    res = client.post("/api/documents", files={"file": ("notes.txt", text.encode("cp1252"), "text/plain")})
    doc = res.json()
    assert res.status_code == 201 and doc["pages"] == 2
    _, pages = extract.extract_pages(stored_path(USER_A, doc["id"], ".txt"), "notes.txt")
    assert [p["text"] for p in pages] == [
        "First paragraph about alpha.\n\nSecond paragraph about beta.",
        "Third paragraph about gamma.",
    ]


def test_markdown_formatting_is_removed():
    md = """---
title: Notes
---
# Enzymes
Enzymes are **proteins** that speed up [reactions](https://example.com).
They are *specific*.

- Lock and key
- Induced fit

| Name | Role |
|------|------|
| Amylase | Starch |

```python
x = 1
```
<!-- a comment -->
"""
    assert extract._markdown_text(md) == "\n\n".join([
        "Enzymes",
        "Enzymes are proteins that speed up reactions. They are specific.",
        "• Lock and key",
        "• Induced fit",
        "Name | Role",
        "Amylase | Starch",
        "x = 1",
    ])


def test_markdown_upload_and_empty_files(client):
    res = client.post("/api/documents", files={"file": ("guide.md", "# Title\n\nSome text.".encode(), "text/markdown")})
    assert res.status_code == 201 and res.json()["chunks"] == 1
    res = client.post("/api/documents", files={"file": ("empty.txt", b"   \n\n", "text/plain")})
    assert res.status_code == 422
    assert "Scanned" not in res.json()["detail"]
    assert rag.chunk_count() == 1


def test_utf16_text():
    assert extract.looks_valid("a.txt", "héllo".encode("utf-16"))
